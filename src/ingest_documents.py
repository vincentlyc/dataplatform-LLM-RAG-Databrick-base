from pathlib import Path
import pandas as pd
import zipfile
from docx import Document
from pyspark.sql import functions as F


def read_docx_text(file_path: str) -> str:
    with open(file_path, "rb") as f:
        doc = Document(f)

    parts = []

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    parts.extend(paragraphs)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def ingest_docx_folder_to_bronze(spark, base_path: str, target_table: str):
    rows = []

    for p in Path(base_path).glob("*.docx"):
        if not p.is_file():
            continue
        if not zipfile.is_zipfile(str(p)):
            continue

        text = read_docx_text(str(p))
        if not text.strip():
            continue

        rows.append(
            {
                "doc_id": p.stem,
                "doc_title": p.stem,
                "file_name": p.name,
                "source_path": str(p),
                "source_url": None,
                "content_raw": text,
            }
        )

    if not rows:
        raise ValueError("No valid docx files found")

    pdf = pd.DataFrame(rows)
    sdf = spark.createDataFrame(pdf)

    (
        sdf.withColumn("ingest_time", F.current_timestamp())
        .write.format("delta")
        .mode("overwrite")
        .option("overwriteSchema", "true")
        .saveAsTable(target_table)
    )
